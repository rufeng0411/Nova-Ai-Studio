#!/usr/bin/env python3
"""Extract slide text from PPTX into JSON for document-import."""
from __future__ import annotations

import argparse
import json
import sys

from pptx import Presentation


def extract_pptx_text(path: str) -> list[dict[str, str]]:
    prs = Presentation(path)
    slides: list[dict[str, str]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = str(shape.text or "").strip()
            if text:
                parts.append(text)
        slides.append({"slide": str(idx), "text": "\n".join(parts)})
    return slides


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract PPTX text to JSON")
    parser.add_argument("--input", required=True)
    args = parser.parse_args()
    slides = extract_pptx_text(args.input)
    sys.stdout.write(json.dumps({"slides": slides}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1)
