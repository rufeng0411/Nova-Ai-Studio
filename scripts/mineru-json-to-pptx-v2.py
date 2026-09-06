#!/usr/bin/env python3
# PD-SAAS-FORK: MinerU JSON → layered editable PPTX (inpainted background + text boxes).
"""MinerU JSON → layered editable PPTX: text removed from background, editable text on top."""
from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any

from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR / "lib"))
from pptx_aspect import apply_slide_size, normalize_aspect_ratio  # noqa: E402
from mineru_bbox import bbox_to_pixels, block_text, block_usable_for_layer  # noqa: E402
from slide_layer_inpaint import (  # noqa: E402
    bbox_font_pt,
    estimate_text_color,
    remove_text_from_background,
    save_temp_background,
)

RASTER_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


def _is_raster_image(path: str) -> bool:
    return Path(path).suffix.lower() in RASTER_IMAGE_EXT


def _image_size(path: str) -> tuple[int, int]:
    try:
        from PIL import Image

        with Image.open(path) as im:
            return im.size
    except Exception:
        return 1920, 1080


def _resolve_page_background(
    page_index: int,
    backgrounds: list[str],
    fallback: str | None,
) -> str | None:
    if backgrounds and page_index < len(backgrounds):
        candidate = backgrounds[page_index]
        if _is_raster_image(candidate) and os.path.isfile(candidate):
            return candidate
    if fallback and _is_raster_image(fallback) and os.path.isfile(fallback):
        return fallback
    return None


def _page_size_emu(prs: Presentation) -> tuple[int, int]:
    return int(prs.slide_width), int(prs.slide_height)


def _bbox_to_emu(
    bbox: list[Any],
    slide_w: int,
    slide_h: int,
    ref_w: int,
    ref_h: int,
) -> tuple[Any, Any, Any, Any]:
    x0, y0, x1, y1 = bbox_to_pixels(bbox, ref_w, ref_h)
    left = Emu(int(x0 / ref_w * slide_w))
    top = Emu(int(y0 / ref_h * slide_h))
    width = Emu(max(int((x1 - x0) / ref_w * slide_w), 200000))
    height = Emu(max(int((y1 - y0) / ref_h * slide_h), 150000))
    return left, top, width, height


def _text_from_legacy_block(block: dict[str, Any]) -> str:
    lines = block.get("lines") or [block]
    parts: list[str] = []
    for line in lines:
        if not isinstance(line, dict):
            continue
        for span in line.get("spans") or []:
            if isinstance(span, dict):
                t = str(span.get("content") or span.get("text") or "").strip()
                if t:
                    parts.append(t)
        merged = "".join(parts).strip() or str(line.get("text") or "").strip()
        if merged:
            return merged
    return ""


def _collect_bbox_list(page: dict[str, Any]) -> list[list[Any]]:
    bboxes: list[list[Any]] = []
    for block in page.get("content_list") or []:
        if isinstance(block, dict) and block_usable_for_layer(block):
            bbox = block.get("bbox")
            if isinstance(bbox, list) and len(bbox) >= 4:
                bboxes.append(bbox)
    for block in page.get("para_blocks") or page.get("preproc_blocks") or []:
        if isinstance(block, dict) and _text_from_legacy_block(block):
            bbox = block.get("bbox")
            if isinstance(bbox, list) and len(bbox) >= 4:
                bboxes.append(bbox)
    return bboxes


def _apply_text_style(box, text: str, font_pt: float, rgb: tuple[int, int, int]) -> None:
    box.line.fill.background()
    box.fill.background()
    tf = box.text_frame
    tf.clear()
    tf.text = text
    for paragraph in tf.paragraphs:
        paragraph.font.size = Pt(font_pt)
        paragraph.font.color.rgb = RGBColor(*rgb)
        paragraph.font.bold = False


def _collect_text_spans(
    page: dict[str, Any],
    slide_w: int,
    slide_h: int,
    ref_w: int,
    ref_h: int,
    source_image: Image.Image | None = None,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []

    # MinerU v4 flat content_list items on page
    for block in page.get("content_list") or []:
        if not isinstance(block, dict) or not block_usable_for_layer(block):
            continue
        text = block_text(block)
        if not text:
            continue
        bbox = block.get("bbox")
        if isinstance(bbox, list) and len(bbox) >= 4:
            left, top, width, height = _bbox_to_emu(bbox, slide_w, slide_h, ref_w, ref_h)
            font_pt = bbox_font_pt(bbox, slide_h, ref_h, ref_w)
            rgb = (
                estimate_text_color(source_image, bbox, ref_w, ref_h)
                if source_image is not None
                else (32, 32, 32)
            )
        else:
            left, top, width, height = Inches(0.6), Inches(1), Inches(8), Inches(0.6)
            font_pt, rgb = 14.0, (32, 32, 32)
        items.append(
            {
                "text": text,
                "left": left,
                "top": top,
                "width": width,
                "height": height,
                "font_pt": font_pt,
                "rgb": rgb,
            }
        )

    # Legacy pdf_info para_blocks
    for block in page.get("para_blocks") or page.get("preproc_blocks") or []:
        if not isinstance(block, dict):
            continue
        text = _text_from_legacy_block(block)
        if not text:
            continue
        bbox = block.get("bbox")
        if isinstance(bbox, list) and len(bbox) >= 4:
            left, top, width, height = _bbox_to_emu(bbox, slide_w, slide_h, ref_w, ref_h)
            font_pt = bbox_font_pt(bbox, slide_h, ref_h, ref_w)
            rgb = (
                estimate_text_color(source_image, bbox, ref_w, ref_h)
                if source_image is not None
                else (32, 32, 32)
            )
        else:
            left, top, width, height = Inches(0.6), Inches(1), Inches(8), Inches(0.6)
            font_pt, rgb = 14.0, (32, 32, 32)
        items.append(
            {
                "text": text,
                "left": left,
                "top": top,
                "width": width,
                "height": height,
                "font_pt": font_pt,
                "rgb": rgb,
            }
        )

    return items


def _normalize_pages(payload: dict[str, Any] | list[Any]) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        if not payload:
            return [{}]
        if all(isinstance(item, list) for item in payload):
            return [{"content_list": [b for b in page if isinstance(b, dict)]} for page in payload]
        blocks = [item for item in payload if isinstance(item, dict)]
        if not blocks:
            return [{}]
        if any("page_idx" in item for item in blocks):
            by_page: dict[int, list[dict[str, Any]]] = {}
            for item in blocks:
                idx = int(item.get("page_idx", 0))
                by_page.setdefault(idx, []).append(item)
            return [{"content_list": by_page[k]} for k in sorted(by_page)]
        return [{"content_list": blocks}]

    if isinstance(payload, dict):
        pages = payload.get("pdf_info") or payload.get("pages") or []
        if isinstance(pages, list) and pages:
            return [page for page in pages if isinstance(page, dict)] or [{}]
        content_list = payload.get("content_list")
        if isinstance(content_list, list) and content_list:
            return _normalize_pages(content_list)
    return [{}]


def build_pptx(
    payload: dict[str, Any] | list[Any],
    output_file: str,
    background_image: str | None = None,
    background_images: list[str] | None = None,
    aspect_ratio: str = "16:9",
) -> str:
    prs = Presentation()
    slide_w, slide_h = apply_slide_size(prs, aspect_ratio)
    pages = _normalize_pages(payload)

    backgrounds = [p for p in (background_images or []) if _is_raster_image(p) and os.path.isfile(p)]
    fallback_bg = background_image if _is_raster_image(background_image or "") else None
    temp_backgrounds: list[str] = []

    for page_index, page in enumerate(pages):
        blank = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank)
        page_dict = page if isinstance(page, dict) else {}
        page_bg = _resolve_page_background(page_index, backgrounds, fallback_bg)
        ref_w, ref_h = (1920, 1080)
        source_image: Image.Image | None = None
        layered_bg_path: str | None = None

        if page_bg:
            ref_w, ref_h = _image_size(page_bg)
            with Image.open(page_bg) as raw:
                source_image = raw.convert("RGB")
            text_bboxes = _collect_bbox_list(page_dict)
            if text_bboxes:
                cleaned = remove_text_from_background(page_bg, text_bboxes, ref_w, ref_h)
                layered_bg_path = save_temp_background(cleaned)
                temp_backgrounds.append(layered_bg_path)
                slide.shapes.add_picture(
                    layered_bg_path,
                    0,
                    0,
                    width=prs.slide_width,
                    height=prs.slide_height,
                )
            else:
                slide.shapes.add_picture(page_bg, 0, 0, width=prs.slide_width, height=prs.slide_height)

        spans = _collect_text_spans(page_dict, slide_w, slide_h, ref_w, ref_h, source_image)
        for span in spans:
            box = slide.shapes.add_textbox(span["left"], span["top"], span["width"], span["height"])
            _apply_text_style(box, span["text"], span["font_pt"], span["rgb"])

    if len(prs.slides) == 0:
        prs.slides.add_slide(prs.slide_layouts[6])

    os.makedirs(os.path.dirname(os.path.abspath(output_file)) or ".", exist_ok=True)
    prs.save(output_file)

    for temp_path in temp_backgrounds:
        try:
            os.remove(temp_path)
        except OSError:
            pass

    return output_file


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--aspect-ratio", default="16:9")
    parser.add_argument("--background", default="", help="Single raster slide background (legacy)")
    parser.add_argument("--backgrounds", nargs="*", default=[], help="Per-page raster backgrounds")
    args = parser.parse_args()

    with open(args.json, encoding="utf-8") as f:
        payload = json.load(f)
    bg = args.background.strip() or None
    build_pptx(
        payload,
        args.output,
        bg,
        list(args.backgrounds or []),
        normalize_aspect_ratio(args.aspect_ratio),
    )
    print(args.output)


if __name__ == "__main__":
    main()
