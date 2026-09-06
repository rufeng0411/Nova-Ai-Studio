#!/usr/bin/env python3
"""Document IR JSON → editable PPTX (python-pptx)."""
from __future__ import annotations

import argparse
import json
import os
from pptx import Presentation
from pptx.util import Inches, Pt

def normalize_inline(text: str) -> str:
    value = str(text or "")
    value = value.replace("**", "").replace("__", "").replace("`", "")
    import re
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"\[([^\]]+)\]\[[^\]]*\]", r"\1", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def add_slide_title_body(prs: Presentation, title: str, lines: list[str]) -> None:
    layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title:
        slide.shapes.title.text = normalize_inline(title)
    body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
    if body is not None:
        tf = body.text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = normalize_inline(line)
            p.font.size = Pt(18)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--ir", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    with open(args.ir, encoding="utf-8") as f:
        ir = json.load(f)

    prs = Presentation()
    slide_lines: list[str] = []
    slide_title = ir.get("title") or "Presentation"

    def flush_slide() -> None:
        nonlocal slide_lines, slide_title
        if slide_lines:
            add_slide_title_body(prs, slide_title, slide_lines)
            slide_lines = []

    for block in ir.get("blocks") or []:
        btype = block.get("type")
        if btype == "slideBreak":
            flush_slide()
            continue
        if btype == "heading":
            if slide_lines:
                flush_slide()
            slide_title = block.get("text") or slide_title
            continue
        if btype == "paragraph":
            slide_lines.append(block.get("text") or "")
        elif btype == "blockquote":
            slide_lines.append(block.get("text") or "")
        elif btype == "list":
            for item in block.get("items") or []:
                prefix = f"{item}" if block.get("ordered") else f"• {item}"
                slide_lines.append(prefix)
        elif btype == "table":
            flush_slide()
            blank = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank)
            headers = block.get("headers") or []
            rows = block.get("rows") or []
            text = " | ".join(headers) + "\n" + "\n".join(" | ".join(r) for r in rows)
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12), Inches(5))
            box.text_frame.text = text
        elif btype == "image" and block.get("resolvedPath") and os.path.isfile(block["resolvedPath"]):
            flush_slide()
            blank = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(block["resolvedPath"], Inches(0.5), Inches(0.5), width=Inches(12))
        elif btype == "chartImage" and block.get("resolvedPath") and os.path.isfile(block["resolvedPath"]):
            flush_slide()
            blank = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(block["resolvedPath"], Inches(0.5), Inches(0.5), width=Inches(12))

    flush_slide()
    if len(prs.slides) == 0:
        add_slide_title_body(prs, slide_title, ["(empty)"])

    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    prs.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
