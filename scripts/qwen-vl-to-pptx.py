#!/usr/bin/env python3
"""Qwen-VL fallback: extract slide text and write editable PPTX."""
from __future__ import annotations

import argparse
import base64
import json
import mimetypes
import os
import sys
from typing import Any

import requests
from pptx import Presentation
from pptx.util import Inches, Pt


def encode_image(path: str) -> tuple[str, str]:
    mime, _ = mimetypes.guess_type(path)
    mime = mime or "image/png"
    with open(path, "rb") as handle:
        data = base64.b64encode(handle.read()).decode("ascii")
    return mime, data


def call_qwen_vl(api_key: str, model: str, image_path: str) -> list[str]:
    mime, data = encode_image(image_path)
    url = "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"
    prompt = (
        "Extract all visible text from this image for a presentation. "
        "Return strict JSON: {\"lines\": [\"line1\", \"line2\", ...]} with reading order. "
        "Do not add commentary."
    )
    payload = {
        "model": model,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{data}"}},
                    {"type": "text", "text": prompt},
                ],
            }
        ],
        "response_format": {"type": "json_object"},
    }
    resp = requests.post(
        url,
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json=payload,
        timeout=120,
    )
    resp.raise_for_status()
    body = resp.json()
    content = body["choices"][0]["message"]["content"]
    parsed = json.loads(content)
    lines = parsed.get("lines") if isinstance(parsed, dict) else None
    if not isinstance(lines, list):
        raise RuntimeError("Qwen-VL did not return lines JSON")
    return [str(line).strip() for line in lines if str(line).strip()]


def write_pptx(lines: list[str], output_file: str, background_image: str | None = None) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    if background_image and os.path.isfile(background_image):
        slide.shapes.add_picture(background_image, 0, 0, width=prs.slide_width, height=prs.slide_height)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(12), Inches(6.2))
    tf = box.text_frame
    tf.text = lines[0] if lines else "(empty)"
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
    if lines:
        tf.paragraphs[0].font.size = Pt(18)
    os.makedirs(os.path.dirname(output_file) or ".", exist_ok=True)
    prs.save(output_file)


def main() -> int:
    parser = argparse.ArgumentParser(description="Qwen-VL image to editable PPTX")
    parser.add_argument("--image", required=True, help="Input image path")
    parser.add_argument("--output", required=True, help="Output PPTX path")
    parser.add_argument("--model", default="qwen-vl-max")
    parser.add_argument("--api-key", default=os.environ.get("DASHSCOPE_API_KEY", ""))
    args = parser.parse_args()
    if not args.api_key:
        raise RuntimeError("DASHSCOPE_API_KEY is required for Qwen-VL fallback")
    lines = call_qwen_vl(args.api_key, args.model, args.image)
    write_pptx(lines, args.output, args.image)
    print(f"Wrote Qwen-VL editable PPTX to {args.output}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1)
