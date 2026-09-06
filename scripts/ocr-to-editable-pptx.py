#!/usr/bin/env python3
"""Orchestrate MinerU OCR (cloud/local) or Qwen-VL fallback into editable PPTX."""
from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
import time
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR / "lib"))

from mineru_cloud import MinerUCloudError, run_mineru  # noqa: E402
from mineru_cache import cache_key_for_files, load_cached_mineru, save_cached_mineru  # noqa: E402

RASTER_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


def is_raster_image(path: str) -> bool:
    return Path(path).suffix.lower() in RASTER_IMAGE_EXT


def load_qwen_module():
    import importlib.util

    script = SCRIPT_DIR / "qwen-vl-to-pptx.py"
    spec = importlib.util.spec_from_file_location("qwen_vl_to_pptx", script)
    if spec is None or spec.loader is None:
        raise RuntimeError("Failed to load qwen-vl-to-pptx.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def run_qwen_fallback(image_path: str, output_file: str, model: str, api_key: str) -> None:
    import subprocess

    script = SCRIPT_DIR / "qwen-vl-to-pptx.py"
    cmd = [
        sys.executable,
        str(script),
        "--image",
        image_path,
        "--output",
        output_file,
        "--model",
        model,
        "--api-key",
        api_key,
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or proc.stdout.strip() or "Qwen-VL fallback failed")
    print(proc.stdout.strip())


def run_qwen_multi(image_paths: list[str], output_file: str, model: str, api_key: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    qwen = load_qwen_module()
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for image_path in image_paths:
        lines = qwen.call_qwen_vl(api_key, model, image_path)
        slide = prs.slides.add_slide(blank)
        if is_raster_image(image_path) and os.path.isfile(image_path):
            slide.shapes.add_picture(image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(12), Inches(6.2))
        tf = box.text_frame
        tf.text = lines[0] if lines else "(empty)"
        for line in lines[1:]:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(18)
        if lines:
            tf.paragraphs[0].font.size = Pt(18)
    os.makedirs(os.path.dirname(output_file) or ".", exist_ok=True)
    prs.save(output_file)
    print(f"Wrote Qwen-VL multi-slide PPTX ({len(image_paths)} slides) to {output_file}")


def bundle_raster_images_to_pdf(image_paths: list[str]) -> str:
    import subprocess

    temp_pdf = os.path.join(
        tempfile.gettempdir(),
        f"ocr-bundle-{os.getpid()}-{int(time.time() * 1000)}.pdf",
    )
    script = SCRIPT_DIR / "compose-images-document.mjs"
    proc = subprocess.run(
        [os.environ.get("NODE", "node"), str(script), "--output", temp_pdf, "--images", *image_paths],
        capture_output=True,
        text=True,
        cwd=str(SCRIPT_DIR.parent),
    )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or proc.stdout.strip() or "compose-images-document failed")
    return temp_pdf


def prepare_mineru_inputs(
    input_paths: list[str],
) -> tuple[str, list[str]]:
    """Return MinerU ingest path and per-slide raster backgrounds."""
    raster_paths = [p for p in input_paths if is_raster_image(p)]
    if len(raster_paths) > 1:
        return bundle_raster_images_to_pdf(raster_paths), raster_paths
    if len(raster_paths) == 1:
        return raster_paths[0], raster_paths
    return input_paths[0], []


def _load_builder_module():
    import importlib.util

    builder = SCRIPT_DIR / "mineru-json-to-pptx-v2.py"
    spec = importlib.util.spec_from_file_location("mineru_json_to_pptx_v2", builder)
    if spec is None or spec.loader is None:
        raise RuntimeError("Failed to load mineru-json-to-pptx-v2.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _mineru_strategy() -> str:
    raw = os.environ.get("PILOTDECK_MINERU_STRATEGY", "bundled").strip().lower()
    if raw in ("bundled", "parallel", "sequential"):
        return raw
    return "bundled"


def _mineru_cache_enabled() -> bool:
    return os.environ.get("PILOTDECK_MINERU_CACHE", "1").strip().lower() not in ("0", "false", "no")


def _workspace_root_from_paths(paths: list[str]) -> str:
    for path in paths:
        parts = Path(path).parts
        if "artifacts" in parts:
            idx = parts.index("artifacts")
            if idx > 0:
                return str(Path(*parts[:idx]))
    return str(Path(paths[0]).parent.parent) if paths else os.getcwd()


def _pages_from_payload(mod, payload) -> list[dict]:
    normalized = mod._normalize_pages(payload)
    return normalized if normalized else [{"content_list": []}]


def _run_mineru_per_slide(
    backgrounds: list[str],
    token: str,
    api_url: str,
    mode: str,
    *,
    parallel: bool,
) -> dict:
    mod = _load_builder_module()
    pages: list[dict] = []

    def one_page(bg: str) -> dict:
        payload = run_mineru(bg, token, api_url, mode)
        normalized = _pages_from_payload(mod, payload)
        return normalized[0] if normalized else {"content_list": []}

    if parallel and len(backgrounds) > 1:
        from concurrent.futures import ThreadPoolExecutor, as_completed

        workers = min(4, len(backgrounds))
        ordered: dict[int, dict] = {}
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {pool.submit(one_page, bg): idx for idx, bg in enumerate(backgrounds)}
            for future in as_completed(futures):
                idx = futures[future]
                ordered[idx] = future.result()
        pages = [ordered[i] for i in range(len(backgrounds))]
    else:
        for bg in backgrounds:
            pages.append(one_page(bg))
    return {"pdf_info": pages}


def _resolve_mineru_payload(
    mineru_input: str,
    backgrounds: list[str],
    token: str,
    api_url: str,
    mode: str,
) -> dict | list:
    workspace = _workspace_root_from_paths(backgrounds or [mineru_input])
    cache_paths = backgrounds if backgrounds else [mineru_input]
    cache_key = cache_key_for_files(cache_paths)
    if _mineru_cache_enabled():
        cached = load_cached_mineru(cache_key, workspace)
        if cached is not None:
            print(f"MinerU cache hit ({cache_key[:12]}…)", file=sys.stderr)
            return cached

    strategy = _mineru_strategy()
    mod = _load_builder_module()

    if len(backgrounds) > 1 and strategy == "bundled":
        try:
            payload = run_mineru(mineru_input, token, api_url, mode)
            pages = _pages_from_payload(mod, payload)
            if len(pages) == len(backgrounds):
                payload = {"pdf_info": pages}
            else:
                raise MinerUCloudError(
                    f"bundled parse page count {len(pages)} != {len(backgrounds)}; retrying per-slide"
                )
        except (MinerUCloudError, RuntimeError) as exc:
            print(f"bundled MinerU failed ({exc}); falling back to per-slide", file=sys.stderr)
            payload = _run_mineru_per_slide(
                backgrounds,
                token,
                api_url,
                mode,
                parallel=strategy != "sequential",
            )
    elif len(backgrounds) > 1:
        payload = _run_mineru_per_slide(
            backgrounds,
            token,
            api_url,
            mode,
            parallel=strategy == "parallel",
        )
    else:
        payload = run_mineru(mineru_input, token, api_url, mode)

    if _mineru_cache_enabled():
        save_cached_mineru(cache_key, payload, workspace)
    return payload


def run_mineru_pipeline(
    input_path: str,
    output_file: str,
    token: str,
    api_url: str,
    mode: str,
    background_paths: list[str] | None = None,
    aspect_ratio: str = "16:9",
) -> None:
    import subprocess

    backgrounds = [p for p in (background_paths or []) if is_raster_image(p) and os.path.isfile(p)]
    payload = _resolve_mineru_payload(input_path, backgrounds, token, api_url, mode)

    with tempfile.TemporaryDirectory(prefix="mineru-json-") as tmp:
        json_path = os.path.join(tmp, "mineru.json")
        with open(json_path, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False)
        script_v2 = SCRIPT_DIR / "mineru-json-to-pptx-v2.py"
        script = script_v2 if script_v2.is_file() else SCRIPT_DIR / "mineru-json-to-pptx.py"
        cmd = [
            sys.executable,
            str(script),
            "--json",
            json_path,
            "--output",
            output_file,
            "--aspect-ratio",
            aspect_ratio,
        ]
        if backgrounds:
            cmd.extend(["--backgrounds", *backgrounds])
        elif is_raster_image(input_path):
            cmd.extend(["--background", input_path])
        proc = subprocess.run(cmd, capture_output=True, text=True)
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.strip() or proc.stdout.strip() or "mineru-json-to-pptx failed")
        print(proc.stdout.strip())


def main() -> int:
    parser = argparse.ArgumentParser(description="OCR image/PDF to editable PPTX")
    parser.add_argument("--inputs", nargs="+", required=True, help="Input file paths")
    parser.add_argument("--output", required=True, help="Output PPTX path")
    parser.add_argument("--provider", default=os.environ.get("PILOTDECK_DOCUMENT_OCR_PROVIDER", "mineru"))
    parser.add_argument("--mode", default=os.environ.get("PILOTDECK_DOCUMENT_OCR_MODE", "cloud"))
    parser.add_argument("--api-url", default=os.environ.get("PILOTDECK_DOCUMENT_OCR_API_URL", "https://mineru.net/api/v4"))
    parser.add_argument("--api-key", default=os.environ.get("MINERU_API_TOKEN", ""))
    parser.add_argument("--model", default=os.environ.get("PILOTDECK_DOCUMENT_OCR_MODEL", "qwen-vl-max"))
    parser.add_argument("--fallback", default=os.environ.get("PILOTDECK_DOCUMENT_OCR_FALLBACK", "qwen-vl"))
    parser.add_argument("--dashscope-key", default=os.environ.get("DASHSCOPE_API_KEY", ""))
    parser.add_argument("--aspect-ratio", default=os.environ.get("PILOTDECK_DOCUMENT_COMPOSE_ASPECT_RATIO", "16:9"))
    parser.add_argument(
        "--backgrounds",
        nargs="*",
        default=[],
        help="Per-slide raster backgrounds (when MinerU ingests a bundled PDF)",
    )
    args = parser.parse_args()

    input_paths = args.inputs
    provider = args.provider
    explicit_backgrounds = [p for p in (args.backgrounds or []) if is_raster_image(p)]
    dashscope_key = (
        args.dashscope_key
        or (args.api_key if str(args.api_key).startswith("sk-") else "")
        or os.environ.get("DASHSCOPE_API_KEY", "")
    )

    if provider == "qwen-vl":
        raster_inputs = [p for p in input_paths if is_raster_image(p)]
        if len(raster_inputs) > 1:
            run_qwen_multi(raster_inputs, args.output, args.model, dashscope_key)
        else:
            run_qwen_fallback(raster_inputs[0] if raster_inputs else input_paths[0], args.output, args.model, dashscope_key)
        return 0

    mineru_input, derived_backgrounds = prepare_mineru_inputs(input_paths)
    background_paths = explicit_backgrounds or derived_backgrounds

    try:
        run_mineru_pipeline(
            mineru_input,
            args.output,
            args.api_key,
            args.api_url,
            args.mode,
            background_paths,
            args.aspect_ratio,
        )
        return 0
    except (MinerUCloudError, RuntimeError) as exc:
        if args.fallback != "qwen-vl" or not dashscope_key:
            raise
        print(f"MinerU failed ({exc}); falling back to Qwen-VL", file=sys.stderr)
        if len(background_paths) > 1:
            run_qwen_multi(background_paths, args.output, args.model, dashscope_key)
        else:
            fallback_image = background_paths[0] if background_paths else mineru_input
            run_qwen_fallback(fallback_image, args.output, args.model, dashscope_key)
        return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1)
